"""
Pytest configuration and shared fixtures
"""
import pytest
import sys
import os

# Add backend directory to path
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))


@pytest.fixture
def sample_text():
    """Sample document text for testing"""
    return """
    1. Introduction
    This document describes the system requirements.
    
    2. Requirements
    The system shall implement the following features:
    - Feature A: Description of feature A
    - Feature B: Description of feature B
    
    3. Test Cases
    Test case TC-001: Verify feature A
    Expected result: Feature A works correctly
    Actual result: PASS
    """


@pytest.fixture
def sample_sections():
    """Sample extracted sections for testing"""
    return {
        "1. Introduction": "This document describes the system requirements.",
        "2. Requirements": """The system shall implement the following features:
        - Feature A: Description of feature A
        - Feature B: Description of feature B""",
        "3. Test Cases": """Test case TC-001: Verify feature A
        Expected result: Feature A works correctly
        Actual result: PASS"""
    }


@pytest.fixture
def sample_findings_data():
    """Sample findings data for testing"""
    return [
        {
            "issue": "N/A without justification found in section 2.1",
            "suggestion": "Add justification for why this requirement is not applicable",
            "location": "Section 2.1, paragraph 3"
        },
        {
            "issue": "Vague requirement 'should work properly'",
            "suggestion": "Define specific acceptance criteria",
            "location": "Section 3.2"
        }
    ]


@pytest.fixture
def mock_openai_response():
    """Mock OpenAI API response structure"""
    def _create_response(findings_data):
        from unittest.mock import Mock
        import json
        
        mock_response = Mock()
        mock_choice = Mock()
        mock_message = Mock()
        mock_message.content = json.dumps({"findings": findings_data})
        mock_choice.message = mock_message
        mock_response.choices = [mock_choice]
        
        return mock_response
    
    return _create_response


@pytest.fixture
def temp_docx_file(tmp_path):
    """Create a temporary DOCX file for testing"""
    from docx import Document
    
    doc_path = tmp_path / "test_document.docx"
    doc = Document()
    
    doc.add_heading("Test Document", 0)
    doc.add_paragraph("This is a test document.")
    doc.add_heading("Section 1", level=1)
    doc.add_paragraph("Content of section 1.")
    
    doc.save(str(doc_path))
    
    return doc_path


@pytest.fixture
def mock_progress_callback():
    """Mock progress callback function"""
    from unittest.mock import Mock
    return Mock()


@pytest.fixture(autouse=True)
def reset_environment():
    """Reset any global state before each test"""
    # Add any global state cleanup here if needed
    yield
    # Cleanup after test
    pass
