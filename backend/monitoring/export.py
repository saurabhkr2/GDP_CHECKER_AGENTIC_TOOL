"""
Export utilities for GDP Checker monitoring data.
"""
from __future__ import annotations

import os
from datetime import datetime
from typing import Optional

from .schema import CSV_COLUMNS
from .storage import get_storage, MonitoringStorage


def export_monthly_csv(output_dir: str, period_month: Optional[str] = None, storage: Optional[MonitoringStorage] = None) -> str:
    storage = storage or get_storage()
    if period_month is None:
        period_month = datetime.now().strftime("%B")
    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, f"GDP_Checker_{period_month}.csv")
    return storage.export_to_csv(output_path, period_month)


def export_all_to_csv(output_path: str, storage: Optional[MonitoringStorage] = None) -> str:
    storage = storage or get_storage()
    return storage.export_to_csv(output_path)


def export_to_excel(
    output_path: str,
    period_month: Optional[str] = None,
    storage: Optional[MonitoringStorage] = None,
) -> str:
    """
    Export monitoring data to an Excel workbook (.xlsx).

    Creates two sheets:
    - "Usage": one row per session (all CSV_COLUMNS)
    - "Summary": aggregated totals from get_summary_stats()

    Requires openpyxl. Raises RuntimeError if it is not installed.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError as exc:  # pragma: no cover - depends on environment
        raise RuntimeError(
            "openpyxl is required for Excel export. Install it with 'pip install openpyxl'."
        ) from exc

    storage = storage or get_storage()
    if period_month:
        records = storage.get_records_by_period(period_month)
    else:
        records = storage.get_all_records()

    wb = Workbook()

    # --- Usage sheet ---
    ws = wb.active
    ws.title = "Usage"
    ws.append(CSV_COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for record in records:
        data = record.to_dict()
        ws.append([data.get(col) for col in CSV_COLUMNS])

    # --- Summary sheet ---
    summary = wb.create_sheet("Summary")
    summary.append(["Metric", "Value"])
    for cell in summary[1]:
        cell.font = Font(bold=True)
    stats = storage.get_summary_stats()
    for key, value in stats.items():
        summary.append([key, value])

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb.save(output_path)
    return output_path


def print_summary(storage: Optional[MonitoringStorage] = None) -> None:
    storage = storage or get_storage()
    stats = storage.get_summary_stats()
    print("\n=== GDP Checker Usage Summary ===")
    print(f"Total Sessions:        {stats.get('total_sessions', 0)}")
    print(f"Unique Users:          {stats.get('unique_users', 0)}")
    print(f"Total Findings:        {stats.get('total_findings', 0)}")
    print(f"Comments Approved:     {stats.get('total_comments_approved', 0)}")
    print(f"Total Hours Saved:     {stats.get('total_hours_saved', 0):.1f}")
    print("=" * 35)


def main():
    import argparse
    parser = argparse.ArgumentParser(description="Export GDP Checker monitoring data")
    parser.add_argument("--output", "-o", required=True, help="Output path")
    parser.add_argument("--period", "-p", help="Month to export")
    parser.add_argument("--summary", "-s", action="store_true", help="Print summary")
    parser.add_argument("--db", help="Database path")
    args = parser.parse_args()
    
    storage = get_storage(args.db) if args.db else get_storage()
    
    if args.summary:
        print_summary(storage)
        return 0
    
    try:
        if args.output.lower().endswith(".xlsx"):
            output = export_to_excel(args.output, args.period, storage)
        elif args.period:
            if os.path.isdir(args.output):
                output = export_monthly_csv(args.output, args.period, storage)
            else:
                output = storage.export_to_csv(args.output, args.period)
        else:
            output = export_all_to_csv(args.output, storage)
        print(f"✅ Exported to: {output}")
    except Exception as e:
        print(f"❌ Error: {e}")
        return 1
    return 0


if __name__ == "__main__":
    exit(main())
