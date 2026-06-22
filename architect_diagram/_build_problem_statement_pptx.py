"""
Generates problem_statement.pptx for the DHF / GDP Checker initiative.

Run from repo root:
    python architect_diagram/_build_problem_statement_pptx.py
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

OUT = Path(__file__).parent / "problem_statement.pptx"

# Palette (consistent with architecture deck)
AZURE_BLUE   = RGBColor(0x00, 0x78, 0xD4)
AZURE_DARK   = RGBColor(0x0B, 0x3D, 0x91)
AZURE_LIGHT  = RGBColor(0xEA, 0xF3, 0xFF)
RED          = RGBColor(0xC0, 0x39, 0x2B)
RED_LIGHT    = RGBColor(0xFB, 0xE9, 0xE7)
GREEN        = RGBColor(0x2E, 0x7D, 0x32)
GREEN_LIGHT  = RGBColor(0xE8, 0xF5, 0xE9)
AMBER        = RGBColor(0xE0, 0x8E, 0x00)
AMBER_LIGHT  = RGBColor(0xFF, 0xF6, 0xE5)
GREY_TEXT    = RGBColor(0x33, 0x33, 0x33)
GREY_LIGHT   = RGBColor(0xF2, 0xF2, 0xF2)
WHITE        = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_text(slide, x, y, w, h, text, *, size=14, bold=False, color=GREY_TEXT,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = color
    return tb


def add_box(slide, x, y, w, h, text, *, fill=AZURE_BLUE, border=AZURE_DARK,
            text_color=WHITE, size=12, bold=True, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
            align=PP_ALIGN.CENTER):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.color.rgb = border; s.line.width = Pt(1.25)
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.1)
    tf.margin_top = tf.margin_bottom = Inches(0.05)
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = line
        r.font.size = Pt(size if i == 0 else max(size - 2, 9))
        r.font.bold = bold if i == 0 else False
        r.font.color.rgb = text_color
    return s


def slide_header(slide, title, subtitle=None):
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                                  prs.slide_width, Inches(0.55))
    band.fill.solid(); band.fill.fore_color.rgb = AZURE_DARK
    band.line.fill.background()
    add_text(slide, Inches(0.3), Inches(0.05), Inches(13), Inches(0.45),
             title, size=22, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    if subtitle:
        add_text(slide, Inches(0.3), Inches(0.6), Inches(13), Inches(0.3),
                 subtitle, size=12, color=GREY_TEXT)


# ============================================================
# SLIDE 1 — Title
# ============================================================
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                        prs.slide_width, prs.slide_height)
bg.fill.solid(); bg.fill.fore_color.rgb = AZURE_DARK; bg.line.fill.background()

add_text(s, Inches(0.5), Inches(2.3), Inches(12.3), Inches(0.6),
         "Problem Statement", size=22, color=AZURE_LIGHT, align=PP_ALIGN.CENTER)
add_text(s, Inches(0.5), Inches(2.9), Inches(12.3), Inches(1.2),
         "Automating DHF / GDP Document Reviews",
         size=42, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
add_text(s, Inches(0.5), Inches(4.2), Inches(12.3), Inches(0.5),
         "From days of manual checking to minutes of AI-assisted, human-approved review",
         size=18, color=AZURE_LIGHT, align=PP_ALIGN.CENTER)
add_text(s, Inches(0.5), Inches(6.6), Inches(12.3), Inches(0.4),
         "Stakeholder Briefing · v1.0",
         size=12, color=AZURE_LIGHT, align=PP_ALIGN.CENTER)

# ============================================================
# SLIDE 2 — Context
# ============================================================
s = prs.slides.add_slide(BLANK)
slide_header(s, "Context", "Why does this problem exist?")
ctx = [
    ("Regulated domain",
     "Medical-device R&D produces Design History Files (DHF) and Good Documentation Practice (GDP) artefacts that must comply with strict templates and writing rules (e.g. 21 CFR 820.30)."),
    ("Manual review today",
     "Reviewers spend hours per document checking placeholders, acronyms, date formats, requirement testability, traceability, template structure, etc., across DOCX files of 50–200 pages."),
    ("High cost of mistakes",
     "Missed defects propagate downstream to V&V, regulatory submissions, and audits — leading to rework, delayed launches, and audit findings."),
    ("Scale",
     "Hundreds of documents per program, multiple programs in flight; expert reviewers are the bottleneck."),
]
y = Inches(1.2)
for title, desc in ctx:
    add_box(s, Inches(0.5), y, Inches(3.2), Inches(0.85), title,
            fill=AZURE_BLUE, border=AZURE_DARK, size=14, align=PP_ALIGN.LEFT)
    # Left-align the body
    tb = s.shapes.add_textbox(Inches(3.9), y + Inches(0.1), Inches(9.1), Inches(0.75))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
    r = p.add_run(); r.text = desc
    r.font.size = Pt(13); r.font.color.rgb = GREY_TEXT
    y += Inches(1.0)

# ============================================================
# SLIDE 3 — Problem Statement (the centerpiece)
# ============================================================
s = prs.slides.add_slide(BLANK)
slide_header(s, "Problem Statement")

# Large quote-style card
card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                          Inches(0.7), Inches(1.2),
                          Inches(11.9), Inches(2.6))
card.fill.solid(); card.fill.fore_color.rgb = AZURE_LIGHT
card.line.color.rgb = AZURE_DARK; card.line.width = Pt(2.0)

quote = (
    "Reviewing DHF / GDP Word documents for template and writing-quality "
    "compliance is slow, inconsistent and expert-dependent. Existing AI tools "
    "either miss domain-specific issues, generate too much low-value noise, "
    "or — most critically — risk inserting AI-authored content into regulated "
    "documents without human approval, which is unacceptable.\n\n"
    "We need an automated quality-review system that is accurate, "
    "context-aware, configurable, and gated by mandatory human review — "
    "so reviewers can focus their expertise on the few decisions that matter."
)
tb = s.shapes.add_textbox(Inches(1.0), Inches(1.4), Inches(11.3), Inches(2.3))
tf = tb.text_frame; tf.word_wrap = True
for i, line in enumerate(quote.split("\n")):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run(); r.text = line
    r.font.size = Pt(15); r.font.color.rgb = AZURE_DARK
    if i == 0: r.font.bold = True

# Three target outcomes below
add_text(s, Inches(0.7), Inches(4.0), Inches(11.9), Inches(0.4),
         "What success looks like", size=16, bold=True, color=AZURE_DARK)

cols = [
    ("Faster",       "Cut review turnaround from hours per document to under 30 minutes."),
    ("Higher-signal", "≤ 5 high-value comments per section; no grammar / noise spam."),
    ("Compliant",    "Zero AI text in the DOCX without explicit human approval (HITL gate)."),
]
x = Inches(0.7); w = Inches(3.9); gap = Inches(0.15)
for title, body in cols:
    add_box(s, x, Inches(4.5), w, Inches(0.6), title,
            fill=GREEN, border=GREEN, size=14)
    add_box(s, x, Inches(5.15), w, Inches(1.8), body,
            fill=GREEN_LIGHT, border=GREEN, text_color=GREEN, size=12, bold=False,
            align=PP_ALIGN.LEFT)
    x += w + gap

# ============================================================
# SLIDE 4 — Stakeholder Pain Points (themed)
# ============================================================
s = prs.slides.add_slide(BLANK)
slide_header(s, "Stakeholder Pain Points",
             "Consolidated from reviewer interviews")

themes = [
    ("Tool reliability", RED, [
        "Tool not accessible / server unreachable",
        "Fails on Excel files; PDF output inconsistent with DOCX",
        "Cannot read images / charts; flags whole section empty when one cell is blank",
    ]),
    ("Comment quality", AMBER, [
        "Too many low-value comments; grammar dominates",
        "Repetitive / duplicate comments unrelated to context",
        "False positives: figures flagged missing, headings flagged empty, wrong date format",
        "Design docs commented as if they were verification docs",
    ]),
    ("Template & context", RED, [
        "Outdated templates not recognised",
        "Header / Footer content not parsed",
        "TOC, revision history, authorship, references mis-interpreted",
        "Glossary / Terminology section ignored for acronym checks",
    ]),
    ("Capability gaps", AMBER, [
        "No Excel (PLDM, RFS, DMR) support",
        "No Windchill (WC) cross-reference",
        "No document-type awareness; same checks run on every doc type",
        "No user customisation (which checks to run)",
        "Single workflow only — no specialised agents",
    ]),
]
y = Inches(1.1); col_w = Inches(6.25); col_h = Inches(2.85)
positions = [(Inches(0.4), y), (Inches(6.85), y),
             (Inches(0.4), y + col_h + Inches(0.2)),
             (Inches(6.85), y + col_h + Inches(0.2))]
for (title, color, items), (x, ty) in zip(themes, positions):
    add_box(s, x, ty, col_w, Inches(0.5), title,
            fill=color, border=color, size=14)
    body = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                              x, ty + Inches(0.55), col_w, col_h - Inches(0.55))
    body.fill.solid(); body.fill.fore_color.rgb = WHITE
    body.line.color.rgb = color; body.line.width = Pt(1.25)
    tf = body.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.15)
    tf.margin_top = Inches(0.1)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run(); r.text = "• " + item
        r.font.size = Pt(12); r.font.color.rgb = GREY_TEXT
        p.space_after = Pt(4)

# ============================================================
# SLIDE 5 — Why current tools fail
# ============================================================
s = prs.slides.add_slide(BLANK)
slide_header(s, "Why current approaches fail")

rows = [
    ("Pure LLM on flat text", "No structural pre-pass: the model has to *find* the Glossary, the test table, the revision row — it often misses them."),
    ("No domain configuration", "Generic grammar / style checks dominate the output; reviewers tune out and miss the few critical issues."),
    ("No human gate", "AI-authored comments land directly in the DOCX — unacceptable in a regulated environment."),
    ("No deduplication / capping", "Same issue reported by multiple checks; reviewers see 150+ comments per document."),
    ("No document-type awareness", "Design docs get verification checks (testability, traceability) and vice versa."),
    ("Not productionised", "Single-process Flask on a dev box → unreliable access, no audit trail, no scale-out."),
]
y = Inches(1.05)
for name, desc in rows:
    add_box(s, Inches(0.4), y, Inches(3.6), Inches(0.65), name,
            fill=RED, border=RED, size=12)
    add_text(s, Inches(4.15), y + Inches(0.1), Inches(8.9), Inches(0.55),
             desc, size=12, color=GREY_TEXT)
    y += Inches(0.85)

# ============================================================
# SLIDE 6 — Goals & Success Criteria
# ============================================================
s = prs.slides.add_slide(BLANK)
slide_header(s, "Goals & Measurable Success Criteria")

cols = [
    ("Goal", AZURE_DARK),
    ("Today (baseline)", AMBER),
    ("Target", GREEN),
]
# Header row
x = Inches(0.4)
for label, color in cols:
    w = Inches(4.15) if label == "Goal" else Inches(4.2)
    add_box(s, x, Inches(1.05), w, Inches(0.5), label,
            fill=color, border=color, size=13)
    x += w + Inches(0.05)

data = [
    ("Average review time per 80-page DOCX",
     "4–6 hours expert time",
     "≤ 30 min reviewer time"),
    ("Comments shown to reviewer per doc",
     "150–250 (mostly noise)",
     "20–40, ranked by severity"),
    ("AI text inserted without approval",
     "Possible (no gate)",
     "0% (HITL enforced)"),
    ("False-positive rate (empty sections, missing figures)",
     "High (~30%)",
     "< 5%"),
    ("Supported formats",
     "DOCX only",
     "DOCX → Excel → PDF roadmap"),
    ("Auditability (who approved what, when)",
     "None",
     "Full audit trail per decision"),
    ("Availability",
     "Dev box / single process",
     "99.5%+ on Azure Container Apps"),
]
y = Inches(1.6)
for goal, today, target in data:
    add_box(s, Inches(0.4),  y, Inches(4.15), Inches(0.6), goal,
            fill=AZURE_LIGHT, border=AZURE_DARK, text_color=AZURE_DARK,
            size=11, align=PP_ALIGN.LEFT)
    add_box(s, Inches(4.6),  y, Inches(4.2),  Inches(0.6), today,
            fill=AMBER_LIGHT, border=AMBER, text_color=AMBER,
            size=11, align=PP_ALIGN.LEFT)
    add_box(s, Inches(8.85), y, Inches(4.2),  Inches(0.6), target,
            fill=GREEN_LIGHT, border=GREEN, text_color=GREEN,
            size=11, align=PP_ALIGN.LEFT)
    y += Inches(0.7)

# ============================================================
# SLIDE 7 — Scope (in / out)
# ============================================================
s = prs.slides.add_slide(BLANK)
slide_header(s, "Scope", "What we will and will not build")

# In scope
add_box(s, Inches(0.5), Inches(1.1), Inches(6.0), Inches(0.55),
        "In scope (this initiative)",
        fill=GREEN, border=GREEN, size=14, align=PP_ALIGN.LEFT)
in_box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                            Inches(0.5), Inches(1.7), Inches(6.0), Inches(5.3))
in_box.fill.solid(); in_box.fill.fore_color.rgb = WHITE
in_box.line.color.rgb = GREEN; in_box.line.width = Pt(1.5)
in_items = [
    "DOCX quality review (DHF / GDP)",
    "21 quality checks (10 regex + 11 LLM)",
    "Heading-anchored, severity-ranked, deduplicated comments",
    "Per-cell empty detection in tables",
    "Header + footer parsing for authorship / template ID",
    "Human-in-the-loop approval before any comment is applied",
    "Configurable severity threshold and per-section cap",
    "Per-job enable/disable of individual checks (e.g. AI05 grammar)",
    "Annotated DOCX + Markdown summary report",
    "Audit trail of every accept / reject decision",
    "Production hosting on Microsoft Azure (Container Apps, OpenAI, Postgres, Blob)",
]
tf = in_box.text_frame; tf.word_wrap = True
tf.margin_left = Inches(0.2); tf.margin_top = Inches(0.15)
for i, item in enumerate(in_items):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run(); r.text = "✓ " + item
    r.font.size = Pt(12); r.font.color.rgb = GREEN
    p.space_after = Pt(4)

# Out of scope
add_box(s, Inches(6.85), Inches(1.1), Inches(6.0), Inches(0.55),
        "Out of scope (future / separate)",
        fill=RED, border=RED, size=14, align=PP_ALIGN.LEFT)
out_box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                             Inches(6.85), Inches(1.7), Inches(6.0), Inches(5.3))
out_box.fill.solid(); out_box.fill.fore_color.rgb = WHITE
out_box.line.color.rgb = RED; out_box.line.width = Pt(1.5)
out_items = [
    "Excel format support (PLDM, RFS, DMR) — Phase 3",
    "PDF format support — Phase 3",
    "Image / chart understanding (GPT-Vision) — Phase 3",
    "Windchill (WC) cross-reference connector — Phase 3",
    "Multi-tenant / per-tenant prompt library — Phase 4",
    "Specialised agents (Requirements, Risk, Traceability) — Phase 4",
    "Document authoring assistance (drafting new content)",
    "Replacement of QMS review workflow / approvals",
    "Translation / multi-language support",
]
tf = out_box.text_frame; tf.word_wrap = True
tf.margin_left = Inches(0.2); tf.margin_top = Inches(0.15)
for i, item in enumerate(out_items):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run(); r.text = "✗ " + item
    r.font.size = Pt(12); r.font.color.rgb = RED
    p.space_after = Pt(4)

# ============================================================
# SLIDE 8 — Stakeholders
# ============================================================
s = prs.slides.add_slide(BLANK)
slide_header(s, "Stakeholders & Interests")

rows = [
    ("R&D Engineers / Authors", "Faster feedback on their drafts; fewer review iterations."),
    ("QA Reviewers",            "Less manual checking; focus on judgement calls, not noise."),
    ("Program / Project Managers", "Predictable review SLAs; visibility on backlog."),
    ("Regulatory / Compliance", "Demonstrable audit trail; zero unreviewed AI content in final docs."),
    ("IT / Platform",           "Azure-native, managed services, SSO, Key Vault, observability."),
    ("Security & Privacy",      "Tenant data in chosen region; private network paths to LLM."),
    ("Finance / Sponsor",       "Measurable ROI: reviewer hours saved × programs × year."),
]
y = Inches(1.1)
for who, what in rows:
    add_box(s, Inches(0.5), y, Inches(3.5), Inches(0.65), who,
            fill=AZURE_BLUE, border=AZURE_DARK, size=12, align=PP_ALIGN.LEFT)
    add_text(s, Inches(4.15), y + Inches(0.1), Inches(8.9), Inches(0.55),
             what, size=13, color=GREY_TEXT)
    y += Inches(0.8)

# ============================================================
# SLIDE 9 — Phased approach
# ============================================================
s = prs.slides.add_slide(BLANK)
slide_header(s, "Phased Approach", "Incremental delivery, fastest wins first")

phases = [
    ("Phase 1", "Fix the noise (done)",
     "Severity threshold • Dedup • Per-section cap • AI05 off by default • Header parsing • Per-cell empty detection",
     GREEN),
    ("Phase 2", "Context & customisation",
     "Document-type awareness • Structured pre-pass (glossary, tags, tests) • Per-job check selection • Template registry",
     AZURE_BLUE),
    ("Phase 3", "Format coverage",
     "Excel (openpyxl) • PDF (pypdf) • Image / chart understanding • Windchill connector",
     AZURE_BLUE),
    ("Phase 4", "Multi-agent platform",
     "Specialised LangGraph subgraphs (Requirements, Risk, Traceability, Terminology, Structure)",
     AZURE_BLUE),
    ("Phase 5", "Productionisation",
     "Azure Container Apps • Entra ID SSO • Private Endpoints • Key Vault • App Insights • Bicep IaC + GitHub Actions",
     AZURE_BLUE),
]
y = Inches(1.15)
for phase, title, body, color in phases:
    add_box(s, Inches(0.5), y, Inches(1.5), Inches(0.95), phase,
            fill=color, border=color, size=14)
    add_box(s, Inches(2.1), y, Inches(3.5), Inches(0.95), title,
            fill=WHITE, border=color, text_color=color, size=13, align=PP_ALIGN.LEFT)
    add_text(s, Inches(5.75), y + Inches(0.15), Inches(7.3), Inches(0.7),
             body, size=12, color=GREY_TEXT)
    y += Inches(1.05)

# ============================================================
# SLIDE 10 — Call to action
# ============================================================
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                        prs.slide_width, prs.slide_height)
bg.fill.solid(); bg.fill.fore_color.rgb = AZURE_DARK; bg.line.fill.background()

add_text(s, Inches(0.5), Inches(1.5), Inches(12.3), Inches(0.8),
         "What we are asking for", size=30, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER)

asks = [
    ("Endorse the problem statement",
     "Approve scope and success criteria as the basis for this initiative."),
    ("Sponsor the production rollout",
     "Fund hosting (Azure Container Apps, OpenAI, Postgres, Blob, Key Vault, Front Door)."),
    ("Nominate champion reviewers",
     "2–3 expert reviewers per business unit to pilot Phase 1 (noise-reduction) on real DHFs."),
    ("Commit to HITL governance",
     "Confirm that 'no AI text in DOCX without human approval' is a hard organisational rule."),
]
y = Inches(2.7)
for title, desc in asks:
    add_box(s, Inches(1.0), y, Inches(3.5), Inches(0.85), title,
            fill=AZURE_BLUE, border=WHITE, size=13, align=PP_ALIGN.LEFT)
    add_text(s, Inches(4.7), y + Inches(0.1), Inches(7.8), Inches(0.7),
             desc, size=14, color=WHITE)
    y += Inches(1.0)

add_text(s, Inches(0.5), Inches(6.8), Inches(12.3), Inches(0.3),
         "Thank you — questions & discussion",
         size=14, color=AZURE_LIGHT, align=PP_ALIGN.CENTER)

prs.save(OUT)
print(f"Wrote {OUT}")
